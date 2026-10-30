from datetime import date
from pathlib import Path

import pytest

from xhs_collector.collector import (Abort, Collector, Config, Control, DailyCap, Filters, Record, Stop, days_ago,
                                     write_csv, write_xlsx)
from xhs_collector.ui import Node, date_headers, parse_cards, parse_dump

FIX = Path(__file__).parent / "fixtures"
H = 2772
SHOP_A, SHOP_B = "a" * 24, "b" * 24


def test_parse_real_search_cards_carry_shop_tags_and_cart_info():
    cards = parse_cards(parse_dump((FIX / "shoplist.xml").read_text()), H)
    by_shop = {c.shop: c for c in cards}
    assert by_shop["随便考都过的店"].shop_tags == "上新12件"
    c2 = by_shop["不二战的店"]
    assert c2.shop_tags == "3年老店" and "100+人加购" in c2.cart_info and c2.shop_tap
    assert (c2.price, c2.sales_text) == (1.88, "已售617")


def test_parse_real_shop_page_has_no_shop_row_and_truncated_price():
    cards = parse_cards(parse_dump((FIX / "shop1.xml").read_text()), H, shop_row=False)
    assert [c.sales_text for c in cards] == ["已售149", "已售173", "已售718", "已售17"]
    assert all(c.shop is None for c in cards) and cards[-1].price == 1.0


def test_parse_real_new_tab_cards_without_sales_and_date_header():
    nodes = parse_dump((FIX / "shop_new.xml").read_text())
    cards = parse_cards(nodes, H, shop_row=False)
    assert len(cards) == 3 and all(c.sales_text is None for c in cards)
    assert date_headers(nodes) == [(2379, "10-28")]


def test_days_ago_labels():
    today = date(2026, 10, 30)
    assert days_ago("今天", today) == 0 and days_ago("昨天", today) == 1
    assert days_ago("10-28", today) == 2 and days_ago("12-31", today) == 303
    assert days_ago("abc", today) is None and days_ago(None, today) is None


class Fake:
    """模拟搜索页、店铺页（商品/上新两个标签）和商品详情页。"""
    ROW = 700

    def __init__(self, products):
        self.products = products
        self.page, self.shop_id, self.tab, self.sort = "search", None, "goods", None
        self.scroll, self.cur, self.state = 0, None, "list"
        self.stack = []

    def screen_height(self): return H

    def focus(self):
        if self.state == "detail":
            return ("com.xingin.xhs", "GoodsDetailActivityV2")
        return ("com.xingin.xhs", "ShopDetailActivity" if self.page == "shop" else "GlobalSearchActivity")

    def _list(self):
        items = [p for p in self.products if self.page == "search" or p["shop_id"] == self.shop_id]
        if self.page == "shop" and self.tab == "new":
            items = [p for p in items if p.get("date")]
        elif self.page == "shop" and self.sort == "销量":
            items = sorted(items, key=lambda p: -p["n"])
        return items

    def _pos(self, i):
        return 630 * (i % 2) + 52, 700 + (i // 2) * self.ROW - self.scroll

    def dump(self):
        if self.state == "detail":
            p = self.cur
            return [Node("", f"商品名称，{p['title']}", 0, 1414, 1260, 1593), Node(f"¥{p['price']}", "", 56, 1308, 225, 1408),
                    Node(f"已售{p['n']}", "", 1032, 1326, 1204, 1389)]
        nodes = []
        if self.page == "search":
            nodes += [Node("商品", "", 504, 371, 616, 435), Node("综合", "", 84, 529, 183, 584)]
        else:
            nodes += [Node("商品", "", 56, 378, 168, 463), Node("上新", "", 280, 378, 392, 464)]
            nodes += [Node(t, "", 56 + 224 * k, 518, 154 + 224 * k, 588) for k, t in enumerate(["综合", "销量", "新品", "价格"])]
        for i, p in enumerate(self._list()):
            x, y = self._pos(i)
            if y < 600 or y + 400 > H:
                continue
            if self.page == "shop" and self.tab == "new" and i % 2 == 0:
                nodes.append(Node(p["date"], "", 39, y - 90, 148, y - 30))
            nodes.append(Node(p["title"], "", x, y, x + 500, y + 70))
            no_sales = self.page == "shop" and self.tab == "new"
            nodes.append(Node(f"¥{p['price']}" + ("" if no_sales else f"已售{p['n']}"), "", x, y + 160, x + 300, y + 235))
            if self.page == "search":
                nodes.append(Node(p["shop"], "", x + 84, y + 250, x + 400, y + 300))
        return nodes

    def tap(self, x, y):
        if self.state != "list":
            return
        if self.page == "shop":
            for t, rx in (("goods", 56), ("new", 280)):
                if rx <= x <= rx + 112 and 378 <= y <= 463:
                    self.tab, self.scroll = t, 0
                    return
            if y == 553:
                self.sort = {56: "综合", 280: "销量", 504: "新品", 728: "价格"}.get(next((k for k in (56, 280, 504, 728) if k <= x <= k + 98), 0))
                self.scroll = 0
                return
        for i, p in enumerate(self._list()):
            px, py = self._pos(i)
            if px <= x <= px + 500 and py <= y <= py + 70:
                self.state, self.cur = "detail", p
                return
            if self.page == "search" and px + 84 <= x <= px + 400 and py + 250 <= y <= py + 300:
                self.page, self.shop_id, self.tab, self.sort, self.scroll = "shop", p["shop_id"], "goods", None, 0
                return

    def back(self):
        self.state = "list"

    def swipe_up(self, h):
        self.scroll = min(self.scroll + 1400, max(0, (len(self._list()) // 2) * self.ROW - 1500))

    def start_search(self, kw): self.page, self.scroll = "search", 0
    def start_shop(self, sid): self.page, self.shop_id, self.tab, self.sort, self.scroll = "shop", sid, "goods", None, 0
    def top_goods_id(self): return self.cur["id"] if self.state == "detail" else None
    def top_shop_id(self): return self.shop_id if self.page == "shop" else None
    def screenshot(self, path): Path(path).write_bytes(b"png")


def products():
    out = []
    for i in range(1, 9):
        shop = SHOP_A if i <= 5 else SHOP_B
        out.append({"id": f"{i:024x}", "title": f"教资真题资料合集第{i}号", "price": 2.99 + i, "n": 100 * i,
                    "shop": "优学资料屋的店" if shop == SHOP_A else "上岸熊旗舰店", "shop_id": shop,
                    "date": ["今天", "今天", "昨天", "昨天", "01-01", "01-01", None, None][i - 1]})
    return out


def make(tmp_path, fake=None, **cfg):
    fake = fake or Fake(products())
    c = Config(delay=(0, 0), wait_timeout=1.0, out_dir=tmp_path, **cfg)
    return Collector(fake, c, sleep=lambda a, b: None), DailyCap(tmp_path / "state.json", c.daily_cap), fake


def ids(recs): return [r.item_id for r in recs]


def test_shop_mode_collects_whole_shop_with_shop_id_and_source(tmp_path):
    col, cap, _ = make(tmp_path, limit=50)
    recs = col.run_shop(SHOP_A, Filters(), cap)
    assert ids(recs) == [f"{i:024x}" for i in range(1, 6)]
    assert {r.shop_id for r in recs} == {SHOP_A} and recs[0].source == f"店铺:{SHOP_A}"


def test_shop_sort_by_sales_changes_order(tmp_path):
    col, cap, _ = make(tmp_path, limit=3)
    assert ids(col.run_shop(SHOP_A, Filters(), cap, sort="销量")) == [f"{i:024x}" for i in (5, 4, 3)]


def test_new_tab_filters_by_days_and_records_date(tmp_path):
    col, cap, _ = make(tmp_path, limit=50)
    recs = col.run_shop(SHOP_A, Filters(new_days=1), cap, new_tab=True)
    assert [(r.item_id[-1], r.listed_date) for r in recs] == [("1", "今天"), ("2", "今天"), ("3", "昨天"), ("4", "昨天")]


def test_shop_by_name_finds_shop_from_search_cards(tmp_path):
    col, cap, fake = make(tmp_path, limit=50)
    recs = col.run_shop_name("上岸熊旗舰店", Filters(), cap)
    assert ids(recs) == [f"{i:024x}" for i in (6, 7, 8)] and {r.shop_id for r in recs} == {SHOP_B}
    assert recs[0].source == "店铺名:上岸熊旗舰店"


def test_shop_name_not_found_aborts(tmp_path):
    col, cap, _ = make(tmp_path, limit=5)
    with pytest.raises(Abort) as e:
        col.run_shop_name("不存在的店", Filters(), cap)
    assert "没有找到店铺" in str(e.value)


def test_seen_cards_persist_so_rerun_resumes(tmp_path):
    col, cap, _ = make(tmp_path, limit=3)
    first = col.run("教资", Filters(), cap)
    assert len(first) == 3
    col2, cap2, _ = make(tmp_path, limit=3)
    second = col2.run("教资", Filters(), cap2)
    assert ids(second) == [f"{i:024x}" for i in (4, 5, 6)]
    assert DailyCap(tmp_path / "state.json", 100).used == 6


def test_only_with_sales_and_new_items_without_sales(tmp_path):
    col, cap, _ = make(tmp_path, limit=50)
    assert col.run_shop(SHOP_A, Filters(only_with_sales=True), cap, new_tab=True) == []
    assert len(col.run_shop(SHOP_A, Filters(), cap, new_tab=True)) == 5


def test_stop_file_stops_and_pause_file_blocks_until_removed(tmp_path):
    col, cap, _ = make(tmp_path, limit=5)
    (tmp_path / "STOP").write_text("")
    with pytest.raises(Stop):
        col.run("教资", Filters(), cap)
    (tmp_path / "STOP").unlink()
    (tmp_path / "PAUSE").write_text("")
    calls = []
    ctl = Control(tmp_path, sleep=lambda s: (calls.append(s), (tmp_path / "PAUSE").unlink()))
    ctl.check()
    assert calls == [1]


def test_daily_cap_old_state_format_still_loads(tmp_path):
    import json
    from datetime import date as d
    (tmp_path / "state.json").write_text(json.dumps({d.today().isoformat(): 7}))
    assert DailyCap(tmp_path / "state.json", 10).left == 3


def test_csv_and_xlsx_outputs_use_chinese_headers(tmp_path):
    rec = Record("店铺:x", "1" * 24, "标题", 2.99, "已售100", "店", "销量1万+", "3年老店", "100+人加购", "今天", SHOP_A,
                 "2026-10-30T10:00:00")
    write_csv(tmp_path / "a.csv", [rec])
    write_xlsx(tmp_path / "a.xlsx", [rec])
    head = (tmp_path / "a.csv").read_text(encoding="utf-8-sig").splitlines()[0]
    assert head.startswith("来源,商品ID,商品标题") and head.endswith("备注")
    from openpyxl import load_workbook
    ws = load_workbook(tmp_path / "a.xlsx").active
    assert [c.value for c in ws[2]][:4] == ["店铺:x", "1" * 24, "标题", 2.99]
